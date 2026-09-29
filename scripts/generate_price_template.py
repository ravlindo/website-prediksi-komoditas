import json
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Protection, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation


source, target = Path(sys.argv[1]), Path(sys.argv[2])
data = json.loads(source.read_text(encoding="utf-8"))

book = Workbook()
sheet = book.active
sheet.title = "SEMUA DATA"
headers = ["tanggal", "pasar", "komoditas", "satuan", "harga"]
sheet.append(headers)
for row in data["rows"]:
    sheet.append([row[key] for key in headers])

navy, blue, green = "12335F", "276DCE", "20B486"
for cell in sheet[1]:
    cell.fill = PatternFill("solid", fgColor=navy)
    cell.font = Font(color="FFFFFF", bold=True, size=11)
    cell.alignment = Alignment(horizontal="center", vertical="center")
sheet.row_dimensions[1].height = 29
sheet.freeze_panes = "A2"
sheet.auto_filter.ref = f"A1:E{sheet.max_row}"
widths = [15, 27, 43, 16, 18]
for index, width in enumerate(widths, 1):
    sheet.column_dimensions[get_column_letter(index)].width = width

thin = Side(style="thin", color="DCE5EF")
for row in sheet.iter_rows(min_row=2):
    for cell in row:
        cell.border = Border(bottom=thin)
        cell.alignment = Alignment(vertical="center")
    row[4].number_format = '"Rp" #,##0'
    row[4].fill = PatternFill("solid", fgColor="EFF7FF")
sheet.conditional_formatting.add(f"E2:E{sheet.max_row}", ColorScaleRule(start_type="min", start_color="FFF4F4", end_type="max", end_color="DFF7ED"))

reference = book.create_sheet("REFERENSI")
reference.append(["DAFTAR PASAR", "DAFTAR KOMODITAS", "SATUAN"])
for i, market in enumerate(data["markets"], 2):
    reference.cell(i, 1, market)
for i, commodity in enumerate(data["commodities"], 2):
    reference.cell(i, 2, commodity["name"])
    reference.cell(i, 3, commodity["unit"])
reference.sheet_state = "hidden"

market_validation = DataValidation(type="list", formula1=f"=REFERENSI!$A$2:$A${len(data['markets']) + 1}", allow_blank=False)
commodity_validation = DataValidation(type="list", formula1=f"=REFERENSI!$B$2:$B${len(data['commodities']) + 1}", allow_blank=False)
sheet.add_data_validation(market_validation)
sheet.add_data_validation(commodity_validation)
market_validation.add(f"B2:B{sheet.max_row}")
commodity_validation.add(f"C2:C{sheet.max_row}")

guide = book.create_sheet("PETUNJUK", 0)
guide.sheet_view.showGridLines = False
guide.merge_cells("A1:F2")
guide["A1"] = "TEMPLATE IMPORT HARGA KABUPATEN MOJOKERTO"
guide["A1"].fill = PatternFill("solid", fgColor=navy)
guide["A1"].font = Font(color="FFFFFF", bold=True, size=16)
guide["A1"].alignment = Alignment(vertical="center")
instructions = [
    "Isi harga pada sheet SEMUA DATA. Nilai 0 tetap diterima sebagai data kosong.",
    "Jangan mengganti nama lima kolom pada baris pertama.",
    "Format tanggal yang disarankan: YYYY-MM-DD, contoh 2026-08-13.",
    "Nama pasar dan komoditas dapat dipilih dari dropdown agar cocok dengan database.",
    "Setelah selesai, simpan sebagai XLSX lalu unggah melalui menu Import Excel.",
]
for i, text in enumerate(instructions, 4):
    guide.merge_cells(start_row=i, start_column=1, end_row=i, end_column=6)
    guide.cell(i, 1, f"{i - 3:02d}   {text}")
    guide.cell(i, 1).font = Font(size=11, color="263B55")
    guide.cell(i, 1).alignment = Alignment(vertical="center")
    guide.row_dimensions[i].height = 29
guide.column_dimensions["A"].width = 26
for col in "BCDEF":
    guide.column_dimensions[col].width = 16

book.save(target)
